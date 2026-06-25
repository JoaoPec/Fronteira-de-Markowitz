from __future__ import annotations

import argparse
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Font, PatternFill
from scipy.optimize import minimize


TRADING_DAYS = 252
DEFAULT_PATTERN = "*_year1.xlsx"
DEFAULT_PORTFOLIOS = 5_000
DEFAULT_SEED = 42


@dataclass(frozen=True)
class AssetSource:
    ticker: str
    path: Path


@dataclass(frozen=True)
class PortfolioResult:
    name: str
    weights: np.ndarray
    annual_return: float
    annual_risk: float
    sharpe_ratio: float


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Analise carteiras pela teoria moderna de Markowitz.",
    )
    parser.add_argument(
        "--data-dir",
        type=Path,
        default=Path("."),
        help="Diretorio onde estao as planilhas de precos.",
    )
    parser.add_argument(
        "--pattern",
        default=DEFAULT_PATTERN,
        help=f"Padrao para descoberta automatica dos ativos. Padrao: {DEFAULT_PATTERN}",
    )
    parser.add_argument(
        "--asset",
        action="append",
        default=[],
        metavar="TICKER=ARQUIVO.xlsx",
        help="Inclui um ativo manualmente. Pode ser usado varias vezes.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("reports/markowitz_analysis.xlsx"),
        help="Arquivo Excel de saida.",
    )
    parser.add_argument(
        "--chart",
        type=Path,
        default=Path("reports/efficient_frontier.png"),
        help="Imagem PNG da fronteira eficiente.",
    )
    parser.add_argument(
        "--portfolios",
        type=int,
        default=DEFAULT_PORTFOLIOS,
        help=f"Quantidade de carteiras simuladas. Padrao: {DEFAULT_PORTFOLIOS}",
    )
    parser.add_argument(
        "--frontier-points",
        type=int,
        default=60,
        help="Quantidade de pontos otimizados na fronteira eficiente.",
    )
    parser.add_argument(
        "--risk-free-rate",
        type=float,
        default=0.0,
        help="Taxa livre de risco anual em decimal. Ex.: 0.105 para 10,5%% a.a.",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=DEFAULT_SEED,
        help=f"Semente da simulacao aleatoria. Padrao: {DEFAULT_SEED}",
    )
    return parser.parse_args()


def discover_assets(data_dir: Path, pattern: str, manual_assets: list[str]) -> list[AssetSource]:
    if manual_assets:
        assets = []
        for item in manual_assets:
            if "=" not in item:
                raise ValueError(f"Use --asset no formato TICKER=arquivo.xlsx: {item}")

            ticker, file_name = item.split("=", 1)
            path = (data_dir / file_name).resolve()
            assets.append(AssetSource(ticker=ticker.strip().upper(), path=path))
    else:
        assets = [
            AssetSource(ticker=path.stem.split("_")[0].upper(), path=path.resolve())
            for path in sorted(data_dir.glob(pattern))
        ]

    if len(assets) < 2:
        raise ValueError("Informe pelo menos dois ativos para montar uma carteira.")

    duplicated = pd.Series([asset.ticker for asset in assets])
    if duplicated.duplicated().any():
        repeated = ", ".join(sorted(duplicated[duplicated.duplicated()].unique()))
        raise ValueError(f"Tickers duplicados na configuracao: {repeated}")

    missing = [str(asset.path) for asset in assets if not asset.path.exists()]
    if missing:
        raise FileNotFoundError("Arquivos nao encontrados: " + ", ".join(missing))

    return assets


def read_price_history(asset: AssetSource) -> pd.Series:
    data = pd.read_excel(asset.path)
    required_columns = {"Date", "Close"}
    missing = required_columns.difference(data.columns)
    if missing:
        raise ValueError(
            f"{asset.path.name} precisa conter as colunas: {', '.join(sorted(missing))}",
        )

    history = data.loc[:, ["Date", "Close"]].copy()
    history["Date"] = pd.to_datetime(history["Date"], errors="coerce")
    history["Close"] = pd.to_numeric(history["Close"], errors="coerce")
    history = history.dropna().drop_duplicates(subset="Date").sort_values("Date")

    if history.empty:
        raise ValueError(f"{asset.path.name} nao possui precos validos.")

    return history.set_index("Date")["Close"].rename(asset.ticker)


def load_price_matrix(assets: list[AssetSource]) -> pd.DataFrame:
    prices = pd.concat([read_price_history(asset) for asset in assets], axis=1, join="inner")
    prices = prices.dropna(how="any")

    if len(prices) < 3:
        raise ValueError("A amostra precisa ter pelo menos tres datas em comum.")

    return prices


def calculate_daily_returns(prices: pd.DataFrame) -> pd.DataFrame:
    returns = prices.pct_change().dropna(how="any")
    if returns.empty:
        raise ValueError("Nao foi possivel calcular retornos diarios.")
    return returns


def portfolio_metrics(
    weights: np.ndarray,
    expected_daily_returns: pd.Series,
    covariance_matrix: pd.DataFrame,
    risk_free_rate: float,
    trading_days: int = TRADING_DAYS,
) -> tuple[float, float, float]:
    annual_return = float(np.dot(weights, expected_daily_returns) * trading_days)
    annual_risk = float(np.sqrt(weights.T @ covariance_matrix.to_numpy() @ weights) * np.sqrt(trading_days))
    sharpe_ratio = np.nan if annual_risk == 0 else (annual_return - risk_free_rate) / annual_risk
    return annual_return, annual_risk, float(sharpe_ratio)


def make_result(
    name: str,
    weights: np.ndarray,
    expected_daily_returns: pd.Series,
    covariance_matrix: pd.DataFrame,
    risk_free_rate: float,
) -> PortfolioResult:
    annual_return, annual_risk, sharpe_ratio = portfolio_metrics(
        weights,
        expected_daily_returns,
        covariance_matrix,
        risk_free_rate,
    )
    return PortfolioResult(name, weights, annual_return, annual_risk, sharpe_ratio)


def optimize_portfolio(
    expected_daily_returns: pd.Series,
    covariance_matrix: pd.DataFrame,
    risk_free_rate: float,
    objective: str,
    target_return: float | None = None,
) -> np.ndarray:
    asset_count = len(expected_daily_returns)
    initial_weights = np.full(asset_count, 1 / asset_count)
    bounds = tuple((0.0, 1.0) for _ in range(asset_count))
    constraints = [{"type": "eq", "fun": lambda weights: np.sum(weights) - 1}]

    if target_return is not None:
        daily_target = target_return / TRADING_DAYS
        constraints.append(
            {
                "type": "eq",
                "fun": lambda weights: np.dot(weights, expected_daily_returns) - daily_target,
            },
        )

    def annual_risk(weights: np.ndarray) -> float:
        return portfolio_metrics(weights, expected_daily_returns, covariance_matrix, risk_free_rate)[1]

    def negative_sharpe(weights: np.ndarray) -> float:
        sharpe = portfolio_metrics(weights, expected_daily_returns, covariance_matrix, risk_free_rate)[2]
        return 1_000_000 if np.isnan(sharpe) else -sharpe

    objective_function = negative_sharpe if objective == "max_sharpe" else annual_risk

    result = minimize(
        objective_function,
        initial_weights,
        method="SLSQP",
        bounds=bounds,
        constraints=constraints,
        options={"maxiter": 1_000, "ftol": 1e-12},
    )

    if not result.success:
        raise RuntimeError(f"Otimizacao falhou ({objective}): {result.message}")

    return np.clip(result.x, 0, 1)


def simulate_random_portfolios(
    expected_daily_returns: pd.Series,
    covariance_matrix: pd.DataFrame,
    risk_free_rate: float,
    portfolios: int,
    seed: int,
) -> pd.DataFrame:
    rng = np.random.default_rng(seed)
    tickers = list(expected_daily_returns.index)
    weights_matrix = rng.dirichlet(np.ones(len(tickers)), portfolios)
    rows = []

    for index, weights in enumerate(weights_matrix, start=1):
        annual_return, annual_risk, sharpe_ratio = portfolio_metrics(
            weights,
            expected_daily_returns,
            covariance_matrix,
            risk_free_rate,
        )
        row = {
            "Carteira": f"Simulada {index}",
            "Retorno Esperado (a.a.)": annual_return,
            "Risco (a.a.)": annual_risk,
            "Indice de Sharpe": sharpe_ratio,
        }
        row.update({f"Peso {ticker}": weight for ticker, weight in zip(tickers, weights)})
        rows.append(row)

    return pd.DataFrame(rows)


def calculate_efficient_frontier(
    expected_daily_returns: pd.Series,
    covariance_matrix: pd.DataFrame,
    risk_free_rate: float,
    points: int,
) -> pd.DataFrame:
    annual_asset_returns = expected_daily_returns * TRADING_DAYS
    targets = np.linspace(annual_asset_returns.min(), annual_asset_returns.max(), points)
    rows = []

    for target in targets:
        try:
            weights = optimize_portfolio(
                expected_daily_returns,
                covariance_matrix,
                risk_free_rate,
                objective="min_risk",
                target_return=float(target),
            )
            annual_return, annual_risk, sharpe_ratio = portfolio_metrics(
                weights,
                expected_daily_returns,
                covariance_matrix,
                risk_free_rate,
            )
        except RuntimeError:
            continue

        rows.append(
            {
                "Retorno Esperado (a.a.)": annual_return,
                "Risco (a.a.)": annual_risk,
                "Indice de Sharpe": sharpe_ratio,
            },
        )

    return pd.DataFrame(rows)


def build_summary_table(results: list[PortfolioResult], tickers: list[str]) -> pd.DataFrame:
    rows = []
    for result in results:
        row = {
            "Carteira": result.name,
            "Retorno Esperado (a.a.)": result.annual_return,
            "Risco (a.a.)": result.annual_risk,
            "Indice de Sharpe": result.sharpe_ratio,
        }
        row.update({f"Peso {ticker}": weight for ticker, weight in zip(tickers, result.weights)})
        rows.append(row)
    return pd.DataFrame(rows)


def plot_frontier(
    random_portfolios: pd.DataFrame,
    frontier: pd.DataFrame,
    key_results: list[PortfolioResult],
    chart_path: Path,
) -> BytesIO:
    chart_path.parent.mkdir(parents=True, exist_ok=True)

    plt.figure(figsize=(11, 7))
    scatter = plt.scatter(
        random_portfolios["Risco (a.a.)"],
        random_portfolios["Retorno Esperado (a.a.)"],
        c=random_portfolios["Indice de Sharpe"],
        cmap="viridis",
        alpha=0.55,
        s=22,
        edgecolors="none",
    )
    if not frontier.empty:
        plt.plot(
            frontier["Risco (a.a.)"],
            frontier["Retorno Esperado (a.a.)"],
            color="#111827",
            linewidth=2,
            label="Fronteira eficiente otimizada",
        )

    markers = {
        "Carteira de minima variancia": ("#dc2626", "P"),
        "Carteira de maximo Sharpe": ("#2563eb", "*"),
        "Carteira igualmente ponderada": ("#059669", "X"),
    }

    for result in key_results:
        color, marker = markers.get(result.name, ("#111827", "o"))
        plt.scatter(
            result.annual_risk,
            result.annual_return,
            color=color,
            marker=marker,
            s=180,
            edgecolor="white",
            linewidth=1.2,
            label=result.name,
        )

    plt.title("Fronteira Eficiente de Markowitz")
    plt.xlabel("Risco anualizado (desvio-padrao)")
    plt.ylabel("Retorno esperado anualizado")
    plt.colorbar(scatter, label="Indice de Sharpe")
    plt.grid(True, alpha=0.25)
    plt.legend()
    plt.tight_layout()

    image_buffer = BytesIO()
    plt.savefig(chart_path, dpi=160)
    plt.savefig(image_buffer, format="png", dpi=160)
    plt.close()
    image_buffer.seek(0)
    return image_buffer


def format_workbook(writer: pd.ExcelWriter, percent_columns: set[str]) -> None:
    workbook = writer.book
    header_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center")

        for column_cells in worksheet.columns:
            max_length = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
            column_letter = column_cells[0].column_letter
            worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 34)

        headers = {cell.value: cell.column for cell in worksheet[1]}
        for column_name in percent_columns.intersection(headers):
            column_index = headers[column_name]
            for row in worksheet.iter_rows(min_row=2, min_col=column_index, max_col=column_index):
                row[0].number_format = "0.00%"


def export_report(
    output_path: Path,
    chart_buffer: BytesIO,
    prices: pd.DataFrame,
    returns: pd.DataFrame,
    covariance_matrix: pd.DataFrame,
    correlation_matrix: pd.DataFrame,
    random_portfolios: pd.DataFrame,
    frontier: pd.DataFrame,
    summary: pd.DataFrame,
    assets: list[AssetSource],
    risk_free_rate: float,
    seed: int,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    asset_table = pd.DataFrame(
        {
            "Ticker": [asset.ticker for asset in assets],
            "Arquivo": [str(asset.path) for asset in assets],
            "Inicio": [prices.index.min().date()] * len(assets),
            "Fim": [prices.index.max().date()] * len(assets),
            "Observacoes": [len(prices)] * len(assets),
        },
    )
    parameters = pd.DataFrame(
        {
            "Parametro": ["Taxa livre de risco", "Carteiras simuladas", "Seed", "Dias uteis/ano"],
            "Valor": [risk_free_rate, len(random_portfolios), seed, TRADING_DAYS],
        },
    )

    percent_columns = {
        "Retorno Esperado (a.a.)",
        "Risco (a.a.)",
        "Taxa livre de risco",
        *[column for column in summary.columns if column.startswith("Peso ")],
        *[column for column in random_portfolios.columns if column.startswith("Peso ")],
    }

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name="Resumo", index=False)
        asset_table.to_excel(writer, sheet_name="Ativos", index=False)
        parameters.to_excel(writer, sheet_name="Parametros", index=False)
        prices.to_excel(writer, sheet_name="Precos Historicos")
        returns.to_excel(writer, sheet_name="Retornos Diarios")
        (returns - returns.mean()).to_excel(writer, sheet_name="Excesso de Retorno")
        covariance_matrix.to_excel(writer, sheet_name="Matriz Covariancia")
        correlation_matrix.to_excel(writer, sheet_name="Matriz Correlacao")
        frontier.to_excel(writer, sheet_name="Fronteira Eficiente", index=False)
        random_portfolios.to_excel(writer, sheet_name="Carteiras Simuladas", index=False)

        worksheet = writer.book.create_sheet("Grafico")
        chart_buffer.seek(0)
        worksheet.add_image(Image(chart_buffer), "A1")

        format_workbook(writer, percent_columns)


def run_analysis(args: argparse.Namespace) -> tuple[Path, Path, pd.DataFrame]:
    assets = discover_assets(args.data_dir, args.pattern, args.asset)
    prices = load_price_matrix(assets)
    returns = calculate_daily_returns(prices)
    covariance_matrix = returns.cov()
    correlation_matrix = returns.corr()
    expected_daily_returns = returns.mean()
    tickers = list(expected_daily_returns.index)

    equal_weights = np.full(len(tickers), 1 / len(tickers))
    min_variance_weights = optimize_portfolio(
        expected_daily_returns,
        covariance_matrix,
        args.risk_free_rate,
        objective="min_risk",
    )
    max_sharpe_weights = optimize_portfolio(
        expected_daily_returns,
        covariance_matrix,
        args.risk_free_rate,
        objective="max_sharpe",
    )

    key_results = [
        make_result(
            "Carteira igualmente ponderada",
            equal_weights,
            expected_daily_returns,
            covariance_matrix,
            args.risk_free_rate,
        ),
        make_result(
            "Carteira de minima variancia",
            min_variance_weights,
            expected_daily_returns,
            covariance_matrix,
            args.risk_free_rate,
        ),
        make_result(
            "Carteira de maximo Sharpe",
            max_sharpe_weights,
            expected_daily_returns,
            covariance_matrix,
            args.risk_free_rate,
        ),
    ]

    random_portfolios = simulate_random_portfolios(
        expected_daily_returns,
        covariance_matrix,
        args.risk_free_rate,
        args.portfolios,
        args.seed,
    )
    frontier = calculate_efficient_frontier(
        expected_daily_returns,
        covariance_matrix,
        args.risk_free_rate,
        args.frontier_points,
    )
    summary = build_summary_table(key_results, tickers)
    chart_buffer = plot_frontier(random_portfolios, frontier, key_results, args.chart)
    export_report(
        args.output,
        chart_buffer,
        prices,
        returns,
        covariance_matrix,
        correlation_matrix,
        random_portfolios,
        frontier,
        summary,
        assets,
        args.risk_free_rate,
        args.seed,
    )

    return args.output, args.chart, summary


def main() -> None:
    args = parse_args()
    output_path, chart_path, summary = run_analysis(args)
    print("Analise concluida.")
    print(f"Relatorio Excel: {output_path}")
    print(f"Grafico PNG: {chart_path}")
    print("\nResumo das carteiras:")
    print(summary.to_string(index=False))


if __name__ == "__main__":
    main()
